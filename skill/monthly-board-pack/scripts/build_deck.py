"""Build the board deck (8–10 slides: 5 fixed + one per division with a material item) from compute_variances.py output plus commentary text.

Usage:
    python build_deck.py result.json commentary.json out.pptx --entity "ICE Group"

commentary.json keys (all strings / lists of strings, written by Claude in house style):
    headline, divisions {name: [bullet, ...]}, risks [..], next_steps [..]
Numbers on charts and KPI tiles are read from result.json only.
"""
import argparse, json
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

NAVY, BLUE, GREY, INK = RGBColor(0x1E, 0x3A, 0x5F), RGBColor(0x25, 0x63, 0xEB), RGBColor(0x6B, 0x72, 0x80), RGBColor(0x1F, 0x29, 0x37)
RED, GREEN, WHITE, PALE = RGBColor(0xB3, 0x26, 0x1E), RGBColor(0x1D, 0x7A, 0x46), RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0xEF, 0xF6, 0xFF)
W, H = 9144000, 5143500


def m(v):
    return f"{v/1000:,.1f}"


def fa(v):
    return f"MUR {m(v)}m F" if v >= 0 else f"(MUR {m(-v)}m) A"


def box(s, x, y, w, h, fill=None):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(x), Emu(y), Emu(w), Emu(h))
    sh.line.fill.background()
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    return sh


def text(s, x, y, w, h, t, size=14, bold=False, color=INK, fill=None):
    sh = box(s, x, y, w, h, fill)
    tf = sh.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = 0
    lines = t if isinstance(t, list) else [t]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        r = p.add_run(); r.text = ln
        r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(size), bold, color, "Arial"
        p.space_after = Pt(6)
    return sh


def slide(prs, title, sub=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    text(s, 457200, 220000, 8229600, 480000, title, 26, True, NAVY)
    box(s, 457200, 760000, 1200000, 36576, BLUE)
    if sub:
        text(s, 457200, 800000, 8229600, 300000, sub, 12, False, GREY)
    return s


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("result"); ap.add_argument("commentary"); ap.add_argument("out"); ap.add_argument("--entity", default="Group")
    a = ap.parse_args()
    r, c = json.load(open(a.result)), json.load(open(a.commentary))
    k, mon = r["month"]["kpis"], r["month"]
    prs = Presentation(); prs.slide_width, prs.slide_height = W, H

    # 1 Title
    s = prs.slides.add_slide(prs.slide_layouts[6]); box(s, 0, 0, W, H, NAVY); box(s, 0, 0, 73152, H, BLUE)
    text(s, 502920, 1500000, 8000000, 700000, f"{a.entity} — Monthly Board Pack", 34, True, WHITE)
    text(s, 502920, 2250000, 8000000, 400000, f"{r['period']} · Budget vs Actual · MUR m", 18, False, RGBColor(0xCA, 0xDC, 0xFC))

    # 2 Headline KPIs
    s = slide(prs, "Headline", c["headline"])
    for i, name in enumerate(["Revenue", "Total costs", "EBITDA", "Net result"]):
        x = 457200 + i * 2080000; d = k[name]
        box(s, x, 1350000, 1960000, 1700000, PALE)
        text(s, x + 100000, 1400000, 1760000, 300000, name.upper(), 11, True, GREY)
        text(s, x + 100000, 1700000, 1760000, 500000, m(d["actual"]), 26, True, NAVY)
        text(s, x + 100000, 2250000, 1760000, 300000, f"Budget {m(d['budget'])}", 11, False, GREY)
        pct = "n/a" if d["var_pct"] is None else f"{d['var_pct']*100:.1f}%"
        text(s, x + 100000, 2550000, 1760000, 300000, f"{fa(d['variance'])} · {pct}", 11, True, GREEN if d["variance"] >= 0 else RED)
    th = r["thresholds"]
    text(s, 457200, 3300000, 8229600, 300000, f"Materiality: ≥ {th['pct']:.0f}% or ≥ MUR {th['abs']/1000:.1f}m. All figures calculated in code from the source workbook.", 10, False, GREY)

    # 3 EBITDA bridge (stacked bar: invisible base + delta)
    s = slide(prs, "EBITDA bridge", f"Budget {m(k['EBITDA']['budget'])}m → Actual {m(k['EBITDA']['actual'])}m")
    mats = [x for x in mon["material"] if x["Type"] != "Below EBITDA"]
    other = k["EBITDA"]["variance"] - sum(x["Variance"] for x in mats)
    steps = [(x["Line item"].replace(" (IC)", ""), x["Variance"]) for x in sorted(mats, key=lambda x: x["Variance"])]
    if round(other): steps.append(("Other", other))
    cats, base, up, down = ["Budget"], [0], [k["EBITDA"]["budget"] / 1000], [0]
    run = k["EBITDA"]["budget"]
    for n, v in steps:
        lo, hi = sorted([run, run + v]); cats.append(n); base.append(lo / 1000)
        up.append(v / 1000 if v > 0 else 0); down.append(-v / 1000 if v < 0 else 0); run += v
    cats.append("Actual"); base.append(0); up.append(k["EBITDA"]["actual"] / 1000); down.append(0)
    cd = CategoryChartData(); cd.categories = cats
    cd.add_series("base", base); cd.add_series("Favourable / total", up); cd.add_series("Adverse", down)
    ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, Emu(457200), Emu(1150000), Emu(8229600), Emu(3700000), cd).chart
    ch.has_legend = False; ch.plots[0].gap_width = 40
    ch.series[0].format.fill.background()
    ch.series[1].format.fill.solid(); ch.series[1].format.fill.fore_color.rgb = GREEN
    ch.series[2].format.fill.solid(); ch.series[2].format.fill.fore_color.rgb = RED
    for i in (0, len(cats) - 1):
        pt = ch.series[1].points[i]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = NAVY
    ch.category_axis.tick_labels.font.size = Pt(8); ch.value_axis.tick_labels.font.size = Pt(9)

    # Divisional slides: every division with at least one material item, largest |EBITDA variance| first
    mat_divs = {x["Division"] for x in mon["material"]}
    divs = sorted(((n, d) for n, d in mon["by_division_ebitda"].items() if n in mat_divs), key=lambda kv: -abs(kv[1]["Variance"]))
    for name, d in divs:
        s = slide(prs, name, f"EBITDA {m(d['Actual'])}m vs budget {m(d['Budget'])}m · {fa(d['Variance'])}")
        rows = [x for x in mon["material"] if x["Division"] == name]
        tbl_h = 300000 * (len(rows) + 1)
        if rows:
            t = s.shapes.add_table(len(rows) + 1, 4, Emu(457200), Emu(1200000), Emu(4200000), Emu(tbl_h)).table
            for j, hd in enumerate(["Line", "Budget", "Actual", "Var"]):
                t.cell(0, j).text = hd
            for i, x in enumerate(rows, 1):
                for j, v in enumerate([x["Line item"], m(x["Budget"]), m(x["Actual"]), fa(x["Variance"])]):
                    t.cell(i, j).text = v
            for row in t.rows:
                for cell in row.cells:
                    for p in cell.text_frame.paragraphs:
                        for rn in p.runs: rn.font.size = Pt(10)
        text(s, 4850000, 1200000, 3836800, 3500000, c["divisions"].get(name, ["Commentary to be added — see material lines in the table."]), 11)

    # 7 Risks & watch list
    s = slide(prs, "Risks & watch list")
    text(s, 457200, 1200000, 8229600, 3500000, ["• " + x for x in c["risks"]], 14)
    # 8 Next steps
    s = slide(prs, "Decisions & next steps")
    text(s, 457200, 1200000, 8229600, 3500000, [f"{i}. {x}" for i, x in enumerate(c["next_steps"], 1)], 14)
    prs.save(a.out)
    print(f"Saved {a.out} with {len(prs.slides)} slides")


if __name__ == "__main__":
    main()
